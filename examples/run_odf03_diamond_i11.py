#!/usr/bin/env python3
"""ODF-03 Diamond I11 v1.1: full-grid, reproducible provenance test.

Run: python run_odf03_diamond_i11_portable.py --data INPUT_DIR --core CORE_SRC_DIR --output OUTPUT_DIR
Requires: openpyxl and the pre-existing ODF_CORE_v0_1/src/odf_core_v0_1.py.
Does not modify input files or ODF core. Writes only to the caller-specified output directory; source files are read-only.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

def arguments():
    parser = argparse.ArgumentParser(description="ODF-03 Diamond I11 reproducible provenance check")
    parser.add_argument("--data", type=Path, required=True, help="Directory containing four downloaded files")
    parser.add_argument("--core", type=Path, required=True, help="Directory containing odf_core_v0_1.py")
    parser.add_argument("--output", type=Path, required=True, help="New output directory")
    return parser.parse_args()


ZENODO = "https://doi.org/10.5281/zenodo.7871450"
EXPECTED = {
    "578401-mythen_summedSILICON640C.xye": "6D1806FC791FD162036735FBFA2679B445D25F5F5CC3A289AB346A729C16C2F8",
    "578412-mythen_summed.xye": "12317867B92CE27D5A309212A8160ED543AF2BEBC7A7BB4914D96D59D6E7EF6E",
    "578413-mythen_summed.xye": "EB54644F1E09363D35BDA5ECA06A1B2DAA71FEDA9AAD85FBFB9FE15A2BF3CA1D",
    "ProjectM_DiffractionData.xlsx": "703E46C842BA3263180652065A3F6269B2947F85E30C6B80E546B9208A9E4695",
}
SAMPLE_NAMES = ["578412-mythen_summed.xye", "578413-mythen_summed.xye"]
STANDARD = "578401-mythen_summedSILICON640C.xye"


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def need(condition: bool, label: str) -> None:
    if not condition:
        raise ValueError(label)


def read_xye(path: Path) -> dict:
    first = last = None
    n = 0
    sum_i = 0.0
    max_i = float("-inf")
    prev_x = None
    grid = []
    min_step = float("inf")
    max_step = float("-inf")
    with path.open("r", encoding="utf-8-sig") as stream:
        for line_number, line in enumerate(stream, 1):
            if not line.strip():
                continue
            fields = line.split()
            need(len(fields) == 3, f"{path.name}:{line_number}: expected 3 columns")
            x, y, sigma = map(float, fields)
            need(all(map(math.isfinite, (x, y, sigma))), f"{path.name}:{line_number}: nonfinite")
            need(sigma >= 0.0, f"{path.name}:{line_number}: negative uncertainty")
            if prev_x is not None:
                delta = x - prev_x
                need(delta > 0, f"{path.name}:{line_number}: nonincreasing angle")
                min_step = min(min_step, delta)
                max_step = max(max_step, delta)
            grid.append(x)
            if first is None:
                first = x
            last, prev_x = x, x
            n += 1
            sum_i += y
            max_i = max(max_i, y)
    need(n > 1, f"{path.name}: insufficient points")
    return ({"points": n, "axis_first": first, "axis_last": last,
             "axis_step_min": round(min_step, 9), "axis_step_max": round(max_step, 9),
             "mean_intensity": sum_i / n, "max_intensity": max_i}, tuple(grid))


def metadata_by_name(workbook_path: Path) -> dict:
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        need("Asp" in book.sheetnames, "Asp worksheet missing")
        sheet = book["Asp"]
        hits = {}
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if len(row) < 9 or not isinstance(row[1], str):
                continue
            name = row[1].strip()
            if name in SAMPLE_NAMES:
                need(name not in hits, f"duplicate metadata link: {name}")
                need(isinstance(row[0], (int, float)), f"missing condition value for {name}")
                hits[name] = {"worksheet": "Asp", "excel_row": row_number,
                              "experimental_condition_value": float(row[0]),
                              "experimental_condition_units": "not specified in inspected header",
                              "rwp_percent": float(row[2]),
                              "calcite_a_angstrom": float(row[3]),
                              "calcite_c_angstrom": float(row[4]),
                              "calcite_volume_angstrom3": float(row[5]),
                              "vaterite_a_angstrom": float(row[6]) if row[6] is not None else None,
                              "vaterite_c_angstrom": float(row[7]) if row[7] is not None else None,
                              "vaterite_volume_angstrom3": float(row[8]) if row[8] is not None else None}
        need(set(hits) == set(SAMPLE_NAMES), f"metadata missing links: {set(SAMPLE_NAMES)-set(hits)}")
        need(hits[SAMPLE_NAMES[0]]["experimental_condition_value"] == 0.0,
             "unexpected condition for 578412")
        need(abs(hits[SAMPLE_NAMES[1]]["experimental_condition_value"] - 0.1) < 1e-10,
             "unexpected condition for 578413")
        return hits
    finally:
        book.close()


def validate_links(names, meta):
    """Standalone strict referential audit, separate from core origin_check()."""
    return set(names) == set(meta) and all(meta[n]["worksheet"] == "Asp" for n in names)


def main() -> int:
    args = arguments()
    data_dir = args.data.expanduser().resolve()
    core_dir = args.core.expanduser().resolve()
    output_dir = args.output.expanduser().resolve()
    try:
        need(data_dir.is_dir(), f"Input folder absent: {data_dir}")
        need((core_dir / "odf_core_v0_1.py").is_file(), f"Core absent: {core_dir}")
        sys.path.insert(0, str(core_dir))
        from odf_core_v0_1 import ODFEvent, Signal, Tension, Invariant, Interpretation, save_event_graph

        hashes = {}
        sizes = {}
        for name, expected in EXPECTED.items():
            path = data_dir / name
            need(path.is_file(), f"Missing input: {name}")
            hashes[name] = digest(path)
            sizes[name] = path.stat().st_size
            need(hashes[name] == expected, f"SHA256 mismatch: {name}")
        print("INPUT_SHA256: PASS (4/4; baseline recorded in this study, not Zenodo cross-check)")

        parsed = {name: read_xye(data_dir / name) for name in [STANDARD, *SAMPLE_NAMES]}
        metrics = {name: parsed[name][0] for name in parsed}
        grids = {name: parsed[name][1] for name in parsed}
        grid_a, grid_b = (grids[name] for name in SAMPLE_NAMES)
        need(len(grid_a) == len(grid_b), "Different number of angular points")
        need(grid_a == grid_b, "Full angular grids differ")
        # Negative control for full-grid validation: change only an interior coordinate
        # in memory while keeping point count and first/last coordinates untouched.
        corrupted_grid = list(grid_b)
        interior = len(corrupted_grid) // 2
        corrupted_grid[interior] += 0.000001
        grid_negative_caught = tuple(corrupted_grid) != grid_a
        need(grid_negative_caught, "Full-grid negative control failed")
        print(f"XYE_FULL_GRIDS: PASS ({len(grid_a)} coordinates compared)")
        print("NEGATIVE_CONTROL_INTERIOR_GRID: PASS (in-memory change detected)")

        meta = metadata_by_name(data_dir / "ProjectM_DiffractionData.xlsx")
        need(validate_links(SAMPLE_NAMES, meta), "Real metadata referential check failed")
        print("EXCEL_ASP_REFERENCES: PASS (2/2)")

        # An isolated negative control. No source files or original metadata are changed.
        mutated_meta = {k: dict(v) for k, v in meta.items()}
        mutated_meta["INVALID_578413.xye"] = mutated_meta.pop(SAMPLE_NAMES[1])
        negative_caught = not validate_links(SAMPLE_NAMES, mutated_meta)
        need(negative_caught, "Negative control failed: broken linkage accepted")
        print("NEGATIVE_CONTROL_BROKEN_FILENAME: PASS (broken reference rejected)")

        event = ODFEvent(
            event_id="ODF-03-DIAMOND-I11-ASP",
            title="Diamond I11 Project M: source-preserving contrast between Asp conditions",
            origin=(f"Diamond Light Source I11, Project M, proposal EE15723-1, "
                    f"published dataset {ZENODO}; local XYE files and workbook verified by SHA256 baseline."),
            description=("Two XRD diffraction patterns from the Asp worksheet correspond to "
                         "experimental-condition values 0 and 0.1. Silicon file is retained as "
                         "a separately identified standard, with no asserted direct calibration link."))
        event.set_tension(Tension(
            tension_id="T-ASP-0-TO-0.1",
            description="Relationship between changed recorded condition and distinct XRD profile/fit outcomes.",
            observed_change="Asp worksheet condition value 0 -> 0.1; two distinct source files.",
            research_question="Can ODF retain the documented provenance of both profiles and their analytical outcomes?"))
        for sample in SAMPLE_NAMES:
            row = meta[sample]
            m = metrics[sample]
            event.add_signal(Signal(
                signal_id="S-" + sample.split("-")[0],
                signal_type="XRD_XYE_WITH_WORKBOOK_CONTEXT",
                description=(f"Asp worksheet row {row['excel_row']}; condition={row['experimental_condition_value']} "
                             f"(unit unspecified); points={m['points']}; "
                             f"Rwp={row['rwp_percent']}%; "
                             f"SHA256={hashes[sample]}"),
                source=f"{ZENODO} / ProjectM_DiffractionData.xlsx / Asp / {sample}",
                file_ref=sample))
        event.add_signal(Signal(
            signal_id="S-SILICON-578401", signal_type="SEPARATELY_DOCUMENTED_SILICON_STANDARD_XYE",
            description=(f"Silicon standard XYE, {metrics[STANDARD]['points']} points; "
                         f"SHA256={hashes[STANDARD]}; relationship to Asp collection is unverified."),
            source=f"{ZENODO} / {STANDARD}", file_ref=STANDARD))
        event.add_invariant(Invariant(
            invariant_id="INV-REFERENTIAL-ORIGIN",
            name="Source-file identity retained", value="2/2 Asp references matched exact Excel filenames; 4/4 SHA256 baseline",
            description="Workbook relation and file checks were independently validated before event creation."))
        event.add_invariant(Invariant(
            invariant_id="INV-EXPERIMENT-CONTEXT",
            name="Asp provenance context",
            value="Asp: condition 0 -> 0.1; same experimental-group worksheet",
            description="Does not assert repeated observations of the same specimen or causal effect."))
        event.set_interpretation(Interpretation(
            model_name="Provenance-preserving structural comparison (not phase identification)",
            description=("Experimental-condition values and reported fit metrics differ. "
                         "This run audits source relationships and does not infer a new physical phase, "
                         "causal mechanism, or undocumented chronology.")))
        event.add_note("No physical rupture/emergence claimed; rupture list intentionally empty.")
        event.add_note("ODF core origin_check validates required fields only; stricter source and linkage checks are runner-level.")
        event.add_note("The negative control modifies an in-memory copy of the reference key, never original files.")
        origin_check = event.origin_check()
        need(origin_check["complete"], "ODF core origin_check incomplete")
        print("ODF_CORE_ORIGIN_CHECK: PASS (required fields; not integrity proof)")

        manifest = {
            "study": "ODF-03 Diamond I11 Project M", "generated_utc": datetime.now(timezone.utc).isoformat(),
            "data_source": ZENODO, "proposal_id": "EE15723-1", "beamline": "I11",
            "instrument": "Mythen II", "baseline_type": "locally recorded SHA256",
            "files": {name: {"sha256": hashes[name], "size_bytes": sizes[name]} for name in EXPECTED},
            "xye_metrics": metrics, "asp_metadata": meta,
            "audits": {"sha256_baseline": True, "xye_valid": True, "asp_exact_match": True,
                       "negative_broken_reference_detected": negative_caught,
                       "full_angular_grid_equal": True,
                       "negative_interior_grid_detected": grid_negative_caught,
                       "core_required_fields": origin_check["complete"]},
            "limitations": ["No direct silicon-to-Asp calibration relationship established",
                            "Concentration unit unspecified in inspected worksheet header",
                            "Neither physical causation nor beamline chronological steps reconstructed",
                            "SHA256 values not independently compared to Zenodo checksums",
                            "Grid comparison checks numerical coordinate identity, not instrument calibration"]}
        output_dir.mkdir(parents=True, exist_ok=True)
        event.save_json(output_dir / "odf03_event.json")
        event.save_markdown(output_dir / "odf03_event.md")
        save_event_graph(event, output_dir / "odf03_graph.json")
        (output_dir / "odf03_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print("OUTPUT:", output_dir)
        print("ODF03: PASS (structural/provenance scope; NOT independent scientific validation)")
        return 0
    except (OSError, ValueError, ImportError, TypeError) as exc:
        print("ODF03: FAIL:", exc, file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
